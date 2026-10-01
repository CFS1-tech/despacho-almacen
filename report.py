"""Excel descargables (pandas + openpyxl)."""
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from logic import paradas
from schema import BASE_COLS, PEDIDO
from utils import a_fecha, num, telefono_de

NUMERICAS = set(PEDIDO) - {"pedido_gr", "n_orden"} | {"viaje"}


def _formatear(ws, anchos):
    azul = PatternFill("solid", fgColor="1F4FB8")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _a_excel(hojas: list) -> bytes:
    """hojas = [(nombre, DataFrame, anchos, columnas_fecha)]"""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for nombre, df, anchos, fechas in hojas:
            df.to_excel(xw, sheet_name=nombre, index=False)
            ws = xw.sheets[nombre]
            for col in fechas:
                j = list(df.columns).index(col) + 1
                for fila in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    fila[0].number_format = "DD/MM/YYYY"
            _formatear(ws, anchos)
    return buf.getvalue()


def excel_base(df: pd.DataFrame) -> bytes:
    out = pd.DataFrame()
    for k, h in BASE_COLS:
        if k == "fecha_entrega":
            out[h] = df[k].map(a_fecha)
        elif k in NUMERICAS:
            out[h] = df[k].map(lambda v: num(v) if num(v) is not None else "")
        else:
            out[h] = df[k]
    anchos = [12, 22, 28, 9, 9, 8, 16, 18, 16, 16, 36, 26, 14, 22, 14, 12, 8, 8, 8, 8, 8, 8, 9, 40, 14, 13, 14, 18,
              16, 6, 14, 12]
    return _a_excel([("BASE DE ENTREGAS", out, anchos, ["FECHA ENTREGA"])])


def excel_hoja_choferes(rutas: pd.DataFrame, sols: pd.DataFrame) -> bytes:
    s = sols.set_index("id")
    filas = []
    for _, r in rutas.iterrows():
        for n, i in enumerate(paradas(r), 1):
            if i not in s.index:
                continue
            x = s.loc[i]
            filas.append({
                "CÓDIGO RUTA": r["codigo"], "CHOFER": r["chofer"], "VIAJE": num(r["viaje"]) or 1,
                "T  TRANSPORTE": r["t_transporte"], "PLACA": r["placa"], "ORDEN": n,
                "FECHA ENTREGA": a_fecha(x["fecha_entrega"]), "PROCESO": x["proceso"], "CUENTA": x["cuenta"],
                "CLIENTE": x["cliente"], "N°ORDEN": x["n_orden"], "PEDIDO - GR": x["pedido_gr"],
                "HORA DE CITA": x["hora_cita"], "DIRECCION": x["direccion"],
                "CONTACTO": x["contacto"].replace("\n", " / "), "TELEFONO": telefono_de(x["contacto"]),
                "REFERENCIA": x["referencia"], "CANT DE CAJA": num(x["cant_caja"]) or "",
                "PALET'S": num(x["palets"]) or "", "OBSERVACION": x["observacion"], "CÓDIGO DESPACHO": x["codigo"],
            })
    df = pd.DataFrame(filas)
    anchos = [16, 16, 6, 14, 9, 6, 12, 8, 20, 26, 10, 12, 9, 36, 26, 13, 22, 8, 7, 30, 16]
    return _a_excel([("CONDUCTORES", df, anchos, ["FECHA ENTREGA"] if not df.empty else [])])


# ---------------------------------------------------------------- plantilla de importación
ENC_RUTEO = ["FECHA ENTREGA", "CUENTA", "CLIENTE", "HORA DE CITA", "HORA DE LLEGADA", "PROCESO", "ORIGEN", "DESTINO",
             "T  TRANSPORTE", "CHOFER", "DIRECCION", "CONTACTO", "TELEFONO", "REFERENCIA", "PEDIDO - GR", "N°ORDEN",
             "CANT DE CAJA", "CANT INNER", "ACSESORIO", "MAQUINA", "X UND ACC", "PALET'S", "UNIDADES", "OBSERVACION",
             "ENCARGADO", "TRANSPORTE", "REVISIÓN"]


def excel_plantilla(cuentas: list, tipos: list, choferes: list) -> bytes:
    """Plantilla vacía con el formato que lee parser_plantilla.py (hojas MATRICES y RUTEO)."""
    from datetime import datetime, time
    from openpyxl import Workbook

    wb = Workbook()
    azul = PatternFill("solid", fgColor="1F4FB8")
    gris = PatternFill("solid", fgColor="EEF1F0")

    ins = wb.active
    ins.title = "INSTRUCCIONES"
    textos = [
        "CÓMO LLENAR ESTA PLANTILLA",
        "",
        "Hoja MATRICES → columna A (CUENTA): una cuenta por fila, sin repetir. Es la lista del desplegable «Cuenta».",
        "   Las columnas C (T TRANSPORTE), E (PROCESO) y G (CHOFER) son solo de referencia; los tipos y choferes",
        "   se administran en Maestros dentro de la app.",
        "",
        "Hoja RUTEO → una fila por despacho o recepción ya realizada (histórico). No cambies los encabezados",
        "   de la fila 1 ni el orden de las columnas. Las filas de ejemplo (en gris) bórralas antes de subir.",
        "   • FECHA ENTREGA: fecha real de Excel (ej. 28/09/2026). Las filas sin fecha se ignoran.",
        "   • CUENTA: igual que en MATRICES. CLIENTE: destinatario o tienda (ej. 4215.- JOCKEY PLAZA).",
        "   • HORA DE CITA / HORA DE LLEGADA: formato hora (10:00).  PROCESO: OUT o IN.",
        "   • ORIGEN / DESTINO: en OUT el origen es VES SERATRA; en IN el destino es VES SERATRA.",
        "   • CHOFER: si hizo varios viajes el mismo día, escribe «PAUL 2» y la app lo toma como viaje 2.",
        "   • Cantidades (CAJA, INNER, PALET'S, UNIDADES…): solo números.",
        "",
        "Con el histórico la app arma la lista de CLIENTES de cada cuenta (dirección, destino y contacto).",
        "En la app: Maestros → Importar Excel → sube este archivo → «Agregar cuentas y clientes» y «Cargar histórico».",
    ]
    for i, t in enumerate(textos, 1):
        ins.cell(i, 1, t)
    ins["A1"].font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 110

    m = wb.create_sheet("MATRICES")
    for col, enc in ((1, "CUENTA"), (3, "T TRANSPORTE"), (5, "PROCESO"), (7, "CHOFER")):
        c = m.cell(1, col, enc)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), azul
        m.column_dimensions[get_column_letter(col)].width = 26
    for i, v in enumerate(cuentas or ["CUENTA EJEMPLO"], 2):
        m.cell(i, 1, v)
    for i, v in enumerate(tipos, 2):
        m.cell(i, 3, v)
    for i, v in enumerate(["OUT", "IN"], 2):
        m.cell(i, 5, v)
    for i, v in enumerate(choferes, 2):
        m.cell(i, 7, v)

    r = wb.create_sheet("RUTEO")
    r.append(ENC_RUTEO)
    ejemplos = [
        [datetime(2026, 9, 28), cuentas[0] if cuentas else "CUENTA EJEMPLO", "4215.- JOCKEY PLAZA", time(10, 0),
         time(9, 50), "OUT", "VES SERATRA", "SANTIAGO DE SURCO", "AUTO SUPPLY", "MARTIN",
         "Av. Javier Prado Este 4200, Santiago de Surco", "Tienda", "987654321", "Ingreso por andén 2",
         "EG07-3210", "", 3, "", "", "", "", "", 79, "", "CESAR", "SUPPLY", "OK"],
        [datetime(2026, 9, 28), cuentas[0] if cuentas else "CUENTA EJEMPLO", "IMPORTACION", None, time(12, 0), "IN",
         "IMPORTACION", "VES SERATRA", "CLIENTE", "Cliente", "", "", "", "", "T003-1650", "", "", "", "", "", "", 2,
         "", "SE RECIBE 2 PALETAS", "JOSE", "CLIENTE", "OK"],
    ]
    for fila in ejemplos:
        r.append(fila)
    for c in r[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), azul
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for fila in r.iter_rows(min_row=2, max_row=3):
        for c in fila:
            c.fill = gris
    for fila in r.iter_rows(min_row=2, max_row=500, max_col=1):
        fila[0].number_format = "DD/MM/YYYY"
    for fila in r.iter_rows(min_row=2, max_row=500, min_col=4, max_col=5):
        for c in fila:
            c.number_format = "HH:MM"
    anchos = [12, 22, 28, 9, 9, 8, 16, 18, 16, 14, 36, 22, 13, 22, 14, 12, 8, 8, 8, 8, 8, 8, 9, 34, 12, 13, 10]
    for i, w in enumerate(anchos, 1):
        r.column_dimensions[get_column_letter(i)].width = w
    r.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
