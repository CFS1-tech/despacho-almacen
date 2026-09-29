"""Importa tu Excel de trabajo (PROPUESTA_PLANTILLA): hoja MATRICES y hoja RUTEO.

Devuelve DataFrames listos para cargar en CUENTAS, CLIENTES, TIPOS_TRANSPORTE,
CHOFERES e HISTORICO. No se sube ningún dato al repo: todo se lee del archivo
que el administrador sube en la app.
"""
import re
from collections import Counter, defaultdict
from datetime import datetime, time

import openpyxl
import pandas as pd

from schema import PEDIDO, TABLES
from utils import ALMACEN, slug

# Nombres de cuenta escritos de varias formas en el histórico -> nombre de MATRICES
NORM_CUENTA = {
    "MASEF": "MASSEF", "G Y G": "G & G", "G Y R": "G & R", "ATLANTIC": "ATLANTIC BUSSINES",
    "WINE SOLUTION": "THE WINE SOLUTION", "MAC CENTER": "MACCENTER", "SUPERDEPORTE": "SUPER DEPORTE",
    "GRUPO VCB": "VCB",
}

# Tiendas de CASA DE LAS CARCASAS: código -> distrito / ciudad
DISTRITO_TIENDA = {
    "4201": "SANTA ANITA", "4202": "BELLAVISTA", "4203": "AREQUIPA", "4204": "SURQUILLO",
    "4205": "SAN JUAN DE MIRAFLORES", "4206": "CHORRILLOS", "4207": "AREQUIPA", "4208": "VILLA MARIA DEL TRIUNFO",
    "4209": "AREQUIPA", "4211": "AREQUIPA", "4212": "PIURA", "4213": "INDEPENDENCIA", "4214": "ATE",
    "4215": "SANTIAGO DE SURCO", "4216": "SAN JUAN DE MIRAFLORES", "4217": "CUSCO", "4218": "LA MOLINA",
    "4219": "ICA", "4221": "CAJAMARCA", "4222": "JULIACA", "4223": "IQUITOS", "4224": "TRUJILLO",
    "4225": "COMAS", "4226": "CHICLAYO", "4227": "CHICLAYO", "4228": "CAJAMARCA", "4229": "PUCALLPA",
    "4232": "HUANUCO", "4234": "CHICLAYO",
}

TIPOS_DEFECTO = [
    ("VAN SUPPLY", "SUPPLY"), ("AUTO SUPPLY", "SUPPLY"), ("AUTO", "TERCERIZADO"), ("VAN", "TERCERIZADO"),
    ("MINIVAN", "TERCERIZADO"), ("PORTER 10 M3", "TERCERIZADO"), ("PORTER 30 M3", "TERCERIZADO"),
    ("PORTER 90 M3", "TERCERIZADO"), ("FURGON 40 M3", "TERCERIZADO"), ("AEREO", "TERCERIZADO"),
    ("DHL", "TERCERIZADO"), ("AGENCIA", "AGENCIA"), ("OLTURSA", "AGENCIA"), ("CRUZ DEL SUR", "AGENCIA"),
    ("CLIENTE", "CLIENTE"),
]

CHOFERES_DEFECTO = [
    ("MARTIN", "Chofer propio"), ("PAUL", "Chofer propio"), ("JHON", "Chofer propio"),
    ("JOSE CORONEL", "Chofer tercerizado"), ("PAUL CARDENAS", "Chofer tercerizado"),
    ("GUILLERMO OVIEDO", "Chofer tercerizado"), ("JOSE SUSUYA", "Chofer tercerizado"),
    ("JOSE BARRA", "Chofer tercerizado"), ("GEORGE TORRES", "Chofer tercerizado"),
    ("HECTOR QUISPE", "Chofer tercerizado"), ("OLTURSA", "Agencia"), ("CRUZ DEL SUR", "Agencia"), ("DHL", "Agencia"),
]

COLS_RUTEO = [
    "fecha_entrega", "cuenta", "cliente", "hora_cita", "hora_llegada", "proceso", "origen", "destino",
    "t_transporte", "chofer", "direccion", "contacto", "telefono", "referencia", *PEDIDO,
    "observacion", "encargado", "transporte", "revision",
]


def _cl(v):
    if v is None:
        return ""
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, datetime):
        return v.strftime("%H:%M") if v.year < 1901 else v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def df_tipos_defecto():
    return pd.DataFrame([{"id": slug(n), "nombre": n, "categoria": c, "activo": "TRUE"} for n, c in TIPOS_DEFECTO])


def df_choferes_defecto():
    return pd.DataFrame([{"id": "ch-" + slug(n), "nombre": n, "tipo": t, "dni": "", "telefono": "", "activo": "TRUE"}
                         for n, t in CHOFERES_DEFECTO])


def leer_plantilla(archivo) -> dict:
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    res = {}

    # --- MATRICES: columna A = CUENTA
    cuentas = []
    if "MATRICES" in wb.sheetnames:
        for i, r in enumerate(wb["MATRICES"].iter_rows(values_only=True)):
            if i == 0 or not r or not r[0]:
                continue
            n = _cl(r[0]).upper()
            if n and n not in cuentas:
                cuentas.append(n)
    res["cuentas"] = pd.DataFrame([{"id": slug(c), "nombre": c, "activo": "TRUE"} for c in cuentas],
                                  columns=TABLES["CUENTAS"])

    # --- RUTEO: busca la fila de encabezado "FECHA ENTREGA"
    filas = []
    if "RUTEO" in wb.sheetnames:
        encontrado = False
        for r in wb["RUTEO"].iter_rows(values_only=True):
            if not encontrado:
                encontrado = bool(r) and _cl(r[0]).upper() == "FECHA ENTREGA"
                continue
            if not r or not any(r):
                continue
            f = r[0]
            if not isinstance(f, datetime):
                continue
            if f.year < 2010:  # corrige años mal digitados (p. ej. 2006 -> 2026)
                f = f.replace(year=2026)
            fila = {k: _cl(r[i]) if i < len(r) else "" for i, k in enumerate(COLS_RUTEO)}
            fila["fecha_entrega"] = f.strftime("%Y-%m-%d")
            for k in ("cuenta", "proceso", "t_transporte", "transporte", "origen", "destino"):
                fila[k] = fila[k].upper()
            fila["cuenta"] = NORM_CUENTA.get(fila["cuenta"], fila["cuenta"])
            ch = fila["chofer"].upper()
            m = re.match(r"^(.*?)\s*-?\s*(\d)$", ch)
            if m and m.group(1):  # "PAUL 2" -> chofer PAUL, viaje 2
                fila["chofer"], fila["viaje"] = m.group(1).strip(), m.group(2)
            else:
                fila["chofer"], fila["viaje"] = ch, ""
            filas.append(fila)
    res["historico"] = pd.DataFrame(filas, columns=TABLES["HISTORICO"]).fillna("")

    # --- CLIENTES: se deducen del histórico (por cuenta)
    cuentas_set = set(cuentas)
    agr = {}
    nombres = defaultdict(Counter)
    for f in filas:
        cu, c = f["cuenta"], f["cliente"]
        if cu not in cuentas_set or not c or c.upper() in ("IMPORTACION", "IMPORTACIÓN", "CLIENTE"):
            continue
        m = re.match(r"^(\d{4})", c)
        clave = (cu, m.group(1) if m else slug(c))
        d = agr.setdefault(clave, {"cuenta": cu, "direccion": "", "destino": "", "contacto": "", "referencia": "", "usos": 0})
        d["usos"] += 1
        nombres[clave][c] += 1
        if f["direccion"]:
            d["direccion"] = f["direccion"]
        if f["destino"] and f["destino"] not in ("CLIENTE", ALMACEN):
            d["destino"] = f["destino"]
        contacto = "\n".join(x for x in (f["contacto"], f["telefono"]) if x)
        if contacto:
            d["contacto"] = contacto
        if f["referencia"]:
            d["referencia"] = f["referencia"]
    clientes = []
    for (cu, k), d in agr.items():
        nombre = nombres[(cu, k)].most_common(1)[0][0]
        if k in DISTRITO_TIENDA:  # tiendas con código: nombre y distrito uniformes
            nombre = re.sub(r"^(\d{4})\s*[-.]*\s*", r"\1.- ", nombre)
            d["destino"] = DISTRITO_TIENDA[k]
            if re.search(r"LLEVO|TALMA", d["direccion"], re.I) or ("Ica, Per" in d["direccion"] and k != "4219"):
                d["direccion"] = ""
        if d["referencia"] == nombre:
            d["referencia"] = ""
        clientes.append({"id": f"{slug(cu)}__{k if k.isdigit() else slug(nombre)}"[:120], "cuenta_id": slug(cu),
                         "cuenta": cu, "cliente": nombre, **d, "activo": "TRUE"})
    res["clientes"] = pd.DataFrame(clientes, columns=TABLES["CLIENTES"]).drop_duplicates("id")
    return res
